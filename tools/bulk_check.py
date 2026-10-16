# -*- coding: utf-8 -*-
"""
일괄입력용 엑셀 검사 (B-7 3단계 · docs\\CLAUDE_CODE_TASK_261015.md)

사용법:
    python tools\\bulk_check.py files  <세션.json> <일괄입력용 폴더> [--templates <내장본 폴더>]
    python tools\\bulk_check.py templates <내장본 폴더> <③ 최종본 폴더>

files     : 앱이 만든 일괄입력용 7개를 세션 값과 대조한다.
            · 시트 이름(License 시트 없음 포함)
            · 라벨 보존 — 내장본 6개는 레코드 영역 밖이 내장본과 같고, 어류는 라벨·필독 글자가 fish_labels.json 과 같다
            · 레코드 값 = 세션 값(90_TECH_NOTES §4 셀 매핑을 여기서 따로 적어 앱 코드와 독립으로 본다)
            · 수식 개수 = 내장본과 같음(어류는 수식 0)
            · 종 행 매칭(저서 : 양식 종목록의 같은 국명 행) · 양식 외 종 추가 행
            · 수질 확장항목은 관리자 Import 의 라벨 검색 규칙을 흉내 내 '관리자가 읽을 값'을 확인한다
templates : 내장본 6개를 ③ 최종본과 대조한다 — 조사지점 시트만 빠졌는지, 레코드 영역 상수 0(실자료 잔존 검사),
            레코드 영역 밖이 그대로인지, 수식 개수가 같은지.

종료 코드 0 = 통과, 1 = 실패.
※ 원본·출력 파일은 읽기만 한다.
"""
import sys, os, json, glob, re

try:
    import openpyxl
except ImportError:
    print("openpyxl 이 필요합니다:  pip install openpyxl")
    sys.exit(2)

sys.stdout.reconfigure(encoding="utf-8")

# 내장본 : (파일, 입력 시트, 레코드 영역 시작 행, 시작 열) — 1-based. TemplateGen 과 같은 규칙.
TEMPLATES = {
    "benthos": ("benthos.xlsx", "입력", 1, 13, "2602_방류하천_곡교천_저서동물_DB_v1.xlsx"),
    "bird": ("bird.xlsx", "조류", 5, 1, "2602_방류하천_곡교천_조류_DB_v1.xlsx"),
    "mammal": ("mammal.xlsx", "포유류", 5, 1, "2602_방류하천_곡교천_포유류_DB_v1.xlsx"),
    "amphrep": ("amphrep.xlsx", "양서파충류", 5, 1, "2603_생태현황_양서파충류_DB_v1.xlsx"),   # 자료 행 H·I 수식도 지움(CLEARED_FORMULA_COLS)
    "water": ("water.xlsx", "수질", 2, 6, "2602_방류하천_곡교천_수질_DB_v1.xlsx"),
    "habitat": ("habitat.xlsx", "서식 및 수변환경", 2, 6, "2602_방류하천_곡교천_서식수변_DB_v1.xlsx"),
}

# 레코드 영역에서 수식까지 지우는 열(1-based) — 양서파충류 학명(H)·특징(I) VLOOKUP(10-16 (2) 결정)
CLEARED_FORMULA_COLS = {"amphrep": {8, 9}}

fails = []


def check(ok, msg):
    print(("  ✓ " if ok else "  ✗ ") + msg)
    if not ok:
        fails.append(msg)


def is_formula(v):
    return isinstance(v, str) and v.startswith("=")


def norm(v):
    """비교용 정규화 — 숫자는 정수면 정수 글자로, 글자는 앞뒤 공백 제거."""
    if v is None:
        return None
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        f = float(v)
        return str(int(f)) if f.is_integer() else repr(f)
    s = str(v).strip()
    if s == "":
        return None
    try:
        f = float(s)
        if "," not in s:
            return str(int(f)) if f.is_integer() else repr(f)
    except ValueError:
        pass
    return s


def cells(ws):
    out = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                out[(c.row, c.column)] = c.value
    return out


def formula_count(wb):
    return sum(1 for ws in wb.worksheets for v in cells(ws).values() if is_formula(v))


def sheet_names(path):
    return openpyxl.load_workbook(path, read_only=True).sheetnames


# ── templates 모드 ─────────────────────────────────────────────────────────
def check_templates(tdir, src):
    for key, (fname, sheet, r0, c0, srcname) in TEMPLATES.items():
        print(f"\n=== 내장본 {fname} ← {srcname}")
        t = openpyxl.load_workbook(os.path.join(tdir, fname))
        s = openpyxl.load_workbook(os.path.join(src, srcname))
        expect = [n for n in s.sheetnames if n != "조사지점"]
        check(t.sheetnames == expect, f"시트 = 최종본에서 조사지점만 뺀 것 {t.sheetnames}")
        tc, sc = cells(t[sheet]), cells(s[sheet])
        rec = lambda k: k[0] >= r0 and k[1] >= c0
        left = [k for k in tc if rec(k) and not is_formula(tc[k])]
        check(len(left) == 0, f"실자료 잔존 검사 — 레코드 영역(행{r0}·열{c0}~) 상수 {len(left)}개")
        outside = [k for k in set(tc) | set(sc) if not rec(k) and tc.get(k) != sc.get(k)]
        check(len(outside) == 0, f"레코드 영역 밖(라벨·예시·머리글·종목록) 그대로 — 다른 칸 {len(outside)}개 {sorted(outside)[:5]}")
        gone = CLEARED_FORMULA_COLS.get(key, set())
        rf_t = sum(1 for k, v in tc.items() if rec(k) and is_formula(v))
        rf_s = sum(1 for k, v in sc.items() if rec(k) and is_formula(v) and k[1] not in gone)
        check(rf_t == rf_s, f"레코드 영역 수식 {rf_t} = 최종본 {rf_s}" + (f" (최종본의 {sorted(gone)}열 자료 행 수식은 지움)" if gone else ""))
        if gone:
            left_f = [k for k, v in tc.items() if rec(k) and k[1] in gone and is_formula(v)]
            had_f = sum(1 for k, v in sc.items() if rec(k) and k[1] in gone and is_formula(v))
            check(not left_f, f"자료 행의 지울 수식 {had_f}칸 → 남은 {len(left_f)}칸")
        for other in expect:
            if other == sheet:
                continue
            check(cells(t[other]) == cells(s[other]), f"보조 시트 '{other}' 내용 그대로")
        dv_t, dv_s = len(t[sheet].data_validations.dataValidation), len(s[sheet].data_validations.dataValidation)
        mg_t, mg_s = len(t[sheet].merged_cells.ranges), len(s[sheet].merged_cells.ranges)
        print(f"  · 참고 : 데이터 유효성 {dv_t}/{dv_s} · 병합 {mg_t}/{mg_s}")


# ── files 모드 ────────────────────────────────────────────────────────────
def date8(s):
    return s[:10].replace("-", "") if s else None


def meta_rows(meta):
    """전치 양식 공통 1~12행(0-based) — 90_TECH_NOTES §4."""
    return {1: meta.get("YearChsu"), 2: meta.get("Site"), 3: date8(meta.get("SurveyDate")), 4: meta.get("MajorRegion"),
            5: meta.get("MiddleRegion"), 6: meta.get("River"), 7: meta.get("RiverType"), 8: meta.get("Lat"),
            9: meta.get("Lng"), 10: meta.get("Weather"), 11: meta.get("SurveyAgency"), 12: meta.get("Surveyor")}


def check_column(ws, col, expected, label):
    """전치 양식 : {0-based 행: 기대값} 을 한 열(1-based col)에서 대조."""
    bad = []
    for r0, v in expected.items():
        got = ws.cell(r0 + 1, col).value
        if norm(got) != norm(v):
            bad.append((r0 + 1, v, got))
    check(not bad, f"{label} 레코드 값 = 세션 값 ({len(expected)}칸){'' if not bad else ' 다름 ' + str(bad[:5])}")


def file_for(folder, taxon):
    hits = glob.glob(os.path.join(folder, f"*_{taxon}_DB_v1.xlsx"))
    return hits[0] if hits else None


def check_files(session_path, folder, tdir):
    d = json.load(open(session_path, encoding="utf-8-sig"))
    meta, taxa = d["Meta"], d["Taxa"]
    m = meta_rows(meta)

    # 파일명 형식 : 대분류_지점_분류군_DB_v1.xlsx
    for f in glob.glob(os.path.join(folder, "*.xlsx")):
        n = os.path.basename(f)
        check(re.fullmatch(rf"{meta['Project']}_{meta['Site']}_[^_]+_DB_v1\.xlsx", n) is not None, f"파일명 형식 {n}")

    # 어류 ────────────────────────────────────────────────────────────
    f = file_for(folder, "어류")
    print(f"\n=== 어류 {os.path.basename(f)}")
    wb = openpyxl.load_workbook(f)
    check(wb.sheetnames == ["필독", "어류_입력"], f"시트 {wb.sheetnames} (License 없음)")
    ws = wb["어류_입력"]
    fx = taxa["Fish"]
    exp = dict(m)
    exp.update({13: fx.get("CollectionTime"), 14: fx.get("CollectionTool"), 15: fx.get("CollectionFlowState"),
                16: fx.get("RiverChasu"), 17: fx.get("Bedrock"), 18: fx.get("Concrete"), 19: fx.get("Mud"), 20: fx.get("Sand"),
                21: fx.get("FineGravel"), 22: fx.get("Gravel"), 23: fx.get("SmallStone"), 24: fx.get("BigStone"),
                26: fx.get("HabitatRiverType"), 27: fx.get("HabitatFlowState"), 28: fx.get("SurveyUnavailableReason"),
                29: fx.get("Note"), 43: fx.get("DeCount"), 44: fx.get("EfCount"), 45: fx.get("LeCount"), 46: fx.get("TuCount")})
    check_column(ws, 17, exp, "Q열")
    check(formula_count(wb) == 0, f"수식 없음 ({formula_count(wb)}개)")
    labels = json.load(open(os.path.join(tdir, "fish_labels.json"), encoding="utf-8"))
    for sname in ("필독", "어류_입력"):
        diff = [(r, c, v) for r, c, v in labels["sheets"][sname] if norm(wb[sname].cell(r + 1, c + 1).value) != norm(v)]
        check(not diff, f"'{sname}' 라벨 글자 {len(labels['sheets'][sname])}칸 = fish_labels.json{'' if not diff else ' 다름 ' + str(diff[:3])}")
    stray = [(c.row, c.column) for row in ws.iter_rows(min_row=1, max_row=78) for c in row
             if c.value is not None and not (13 <= c.column <= 16 or c.column == 17)]
    check(not stray, f"1~78행에 라벨(M~P)·레코드(Q) 밖 글자 없음 {stray[:5]}")
    sp = [r for r in fx["Rows"] if (r.get("IndividualCount") or 0) > 0]
    got = [(ws.cell(79 + i, 14).value, ws.cell(79 + i, 17).value) for i in range(len(sp))]
    want = [(r["SpeciesKo"], r["IndividualCount"]) for r in sp]
    check([(a, norm(b)) for a, b in got] == [(a, norm(b)) for a, b in want], f"종 행 79행~ 국명(N)·개체수(Q) {len(sp)}종 = 세션")
    check(ws.cell(79 + len(sp), 14).value is None, "종 행 뒤에 남는 행 없음")

    # 저서동물 ─────────────────────────────────────────────────────────
    f = file_for(folder, "저서동물")
    print(f"\n=== 저서동물 {os.path.basename(f)}")
    wb, tw = openpyxl.load_workbook(f), openpyxl.load_workbook(os.path.join(tdir, "benthos.xlsx"))
    check(wb.sheetnames == tw.sheetnames, f"시트 = 내장본 {wb.sheetnames} (License 없음)")
    ws, tws = wb["입력"], tw["입력"]
    bx = taxa["Benthos"]
    exp = dict(m)
    exp.update({13: bx.get("Surbernet30"), 14: bx.get("Surbernet50"), 15: bx.get("Dredge"), 16: bx.get("Ekman"),
                17: bx.get("Watershed"), 18: bx.get("PollutionSource"), 19: bx.get("CanopyCover"), 20: bx.get("Floodplain"),
                21: bx.get("LeveeLeft"), 22: bx.get("LeveeRight"), 23: bx.get("Bedrock"), 24: bx.get("Concrete"), 25: bx.get("Mud"),
                26: bx.get("Sand"), 27: bx.get("FineGravel"), 28: bx.get("Gravel"), 29: bx.get("SmallStone"), 30: bx.get("BigStone"),
                32: bx.get("HabitatRiverType"), 33: bx.get("RiverWidth"), 34: bx.get("WaterWidth"), 35: bx.get("AverageDepth"),
                36: bx.get("AverageVelocity"), 37: bx.get("AirTemperature"), 38: bx.get("WaterTemperature"), 39: bx.get("FlowState"),
                40: bx.get("Transparency"), 41: bx.get("Smell"), 42: bx.get("Note"), 43: bx.get("SurveyUnavailableReason")})
    check_column(ws, 13, exp, "M열")
    check(formula_count(wb) == formula_count(tw), f"수식 개수 {formula_count(wb)} = 내장본 {formula_count(tw)}")
    # 종 행 : 양식 종목록에서 국명(K) 첫 행
    tpl_rows, last = {}, 46
    for r in range(47, tws.max_row + 1):
        name = tws.cell(r, 11).value
        if name:
            tpl_rows.setdefault(str(name).strip(), r)
            last = r
    sp = [r for r in bx["Rows"] if (r.get("IndividualCount") or 0) > 0]
    appended, ok_rows = [], []
    nxt = last + 1
    for r in sp:
        name = r["SpeciesKo"].strip()
        if name in tpl_rows:
            ok_rows.append((tpl_rows[name], name, r["IndividualCount"]))
        else:
            ok_rows.append((nxt, name, r["IndividualCount"]))
            appended.append(name)
            nxt += 1
    bad = [(row, name, cnt, ws.cell(row, 11).value, ws.cell(row, 13).value) for row, name, cnt in ok_rows
           if str(ws.cell(row, 11).value).strip() != name or norm(ws.cell(row, 13).value) != norm(cnt)]
    check(not bad, f"종 행 매칭 {len(sp)}종(양식 종 {len(sp) - len(appended)} · 끝에 추가 {len(appended)}: {appended}){'' if not bad else ' 다름 ' + str(bad[:3])}")
    used = {row for row, _, _ in ok_rows}
    others = [r for r in range(47, ws.max_row + 1) if r not in used and ws.cell(r, 13).value is not None]
    check(not others, f"다른 종 행의 레코드 열(M) 비어 있음 {others[:5]}")
    # 레코드 영역 밖 = 내장본 (덧붙인 종 행 제외)
    tc, oc = cells(tws), cells(ws)
    rec = lambda k: k[1] >= 13 or k[0] > last
    diff = [k for k in set(tc) | set(oc) if not rec(k) and tc.get(k) != oc.get(k)]
    check(not diff, f"라벨 열(I~L)·종목록(A~K) 보존 — 다른 칸 {len(diff)}개 {sorted(diff)[:5]}")
    extra = [k for k in oc if k[1] > 13 and not is_formula(oc[k])]
    check(not extra, f"첫 레코드 열(M) 밖 레코드 열에 값 없음 {extra[:5]}")

    # 조류·포유류·양서파충류 ─────────────────────────────────────────────
    # 학명 : 세션에 없으면 앱이 공식 종목록(species.json)에서 채운다(CLAUDE.md §5) → 같은 규칙으로 기대값을 만든다.
    species_json = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "SDSM_Surveyor_App", "species.json")
    catalog = json.load(open(species_json, encoding="utf-8-sig"))
    latin = {k: {sp["SpeciesKo"]: sp.get("SpeciesEn") for sp in catalog.get(k, [])} for k in ("Bird", "Mammal", "Amphibian")}
    obs = {
        "조류": ("Bird", "조류", "bird.xlsx", lambda r: (r.get("IndividualCount") or 0) > 0,
               {6: "SpeciesKo", 7: "SpeciesEn", 8: "Feature", 15: "Lat", 16: "Lng", 18: "IndividualCount", 19: "MigratoryType",
                20: "Category", 21: "CategoryDetail", 22: "HabitatType", 24: "Note"}, 23),
        "포유류": ("Mammal", "포유류", "mammal.xlsx", lambda r: sum((r.get(f"Trace{i}") or 0) for i in range(1, 13)) > 0,
                {6: "SpeciesKo", 7: "SpeciesEn", 8: "Feature", 15: "Lat", 16: "Lng", 18: "ObservationSite",   # 양식 머리글대로 S열
                 **{18 + i: f"Trace{i}" for i in range(1, 13)}, 32: "Note"}, 31),
        "양서파충류": ("AmphibianReptile", "양서파충류", "amphrep.xlsx", lambda r: sum((r.get(f"Trace{i}") or 0) for i in range(1, 7)) > 0,
                  {6: "SpeciesKo", 7: "SpeciesEn", 8: "Feature", **{8 + i: f"Trace{i}" for i in range(1, 7)}, 21: "Lat", 22: "Lng",
                   24: "MajorCategory", 25: "MiddleCategory", 27: "Note"}, 26),
    }
    for taxon, (key, sheet, tname, keep, cols, surveyor_col) in obs.items():
        f = file_for(folder, taxon)
        print(f"\n=== {taxon} {os.path.basename(f)}")
        wb = openpyxl.load_workbook(f)
        ws = wb[sheet]
        if True:   # 관찰형 3종 모두 내장본 기반
            tw = openpyxl.load_workbook(os.path.join(tdir, tname))
            check(wb.sheetnames == tw.sheetnames, f"시트 = 내장본 {wb.sheetnames} (License 없음)")
            tc, oc = cells(tw[sheet]), cells(ws)
            diff = [k for k in set(tc) | set(oc) if k[0] < 5 and tc.get(k) != oc.get(k)]
            check(not diff, f"머리글 1~4행(예시 포함) 보존 — 다른 칸 {len(diff)}개 {sorted(diff)[:5]}")
            check(formula_count(wb) == formula_count(tw), f"수식 개수 {formula_count(wb)} = 내장본 {formula_count(tw)}")
        rows = [r for r in taxa[key]["Rows"] if keep(r)]
        names = latin["Amphibian" if key == "AmphibianReptile" else key]
        bad = []
        for i, r in enumerate(rows):
            xr = 5 + i
            want = {2: meta.get("YearChsu"), 3: date8(meta.get("SurveyDate")), 4: meta.get("Weather"), 5: meta.get("River"),
                    6: meta.get("Site"), surveyor_col + 1: meta.get("Surveyor")}
            want.update({c + 1: r.get(k) for c, k in cols.items()})
            if r.get("SpeciesEn") is None:
                want[8] = names.get(r["SpeciesKo"])
            for c, v in want.items():
                if norm(ws.cell(xr, c).value) != norm(v):
                    bad.append((xr, c, v, ws.cell(xr, c).value))
        check(not bad, f"자료 행 5행~ {len(rows)}건 = 세션{'' if not bad else ' 다름 ' + str(bad[:4])}")
        check(ws.cell(5 + len(rows), 2).value is None, "자료 행 뒤에 남는 행 없음")
        if key == "Mammal":
            r_used = [5 + i for i in range(len(rows)) if ws.cell(5 + i, 18).value is not None]
            check(not r_used, f"관찰지유형을 R열(종전·관리자 현재 열)에 쓰지 않음 {r_used[:5]}")

    # 수질 ────────────────────────────────────────────────────────────
    f = file_for(folder, "수질")
    print(f"\n=== 수질 {os.path.basename(f)}")
    wb = openpyxl.load_workbook(f)
    ws = wb["수질"]
    tw = openpyxl.load_workbook(os.path.join(tdir, "water.xlsx"))
    check(wb.sheetnames == tw.sheetnames, f"시트 = 내장본 {wb.sheetnames} (License 없음)")
    check(formula_count(wb) == formula_count(tw), f"수식 개수 {formula_count(wb)} = 내장본 {formula_count(tw)}")
    wx = taxa["WaterQuality"]
    exp = dict(m)
    exp.update({13: wx.get("PH"), 14: wx.get("Bod"), 15: wx.get("Cod"), 16: wx.get("Toc"), 17: wx.get("Ss"), 18: wx.get("Dox"),
                19: wx.get("Tp"), 20: wx.get("Ecotoxicity"), 21: wx.get("EColi")})
    check_column(ws, 6, exp, "F열")
    # 관리자 BulkWaterQualityControlViewModel 의 확장항목 검색(행22~, C+D 공백 제거·소문자 Contains, 첫 일치)
    names = {"TN": ["TN", "T-N"], "EC": ["EC", "전기전도도", "ElectricalConductivity", "Electrical"], "Cl": ["Cl", "Cl-", "염소이온"],
             "SO42": ["SO42", "SO42-", "황이온"], "Cu": ["구리", "Cu"], "Zn": ["아연", "Zn"], "Cr": ["크롬", "Cr"], "Turbidity": ["탁도"],
             "Chla": ["클로로필a", "클로로필", "Chla", "Chl-a"], "WaterTemperature": ["수온", "온도"], "WaterDepth": ["수심"],
             "FlowVelocity": ["유속"], "FlowSec": ["초당유량", "m3/sec", "㎥/sec"], "FlowDay": ["일당유량", "m3/day", "㎥/day"]}
    bad = []
    for key, keys in names.items():
        low = [k.lower() for k in keys]
        found = None
        for r in range(23, ws.max_row + 1):
            c2 = str(ws.cell(r, 3).value or "").replace(" ", "").lower()
            c3 = str(ws.cell(r, 4).value or "").replace(" ", "").lower()
            if c2 and any(k in c2 + c3 for k in low):
                found = ws.cell(r, 6).value
                break
        if norm(found) != norm(wx.get(key)):
            bad.append((key, wx.get(key), found))
    check(not bad, f"확장항목 14종 — 관리자 검색으로 읽힐 값 = 세션{'' if not bad else ' 다름 ' + str(bad)}")
    # 덧붙임 : 양식에 없는 항목은 '지속 삽입 가능' 행 아래 빈 행에 C열 라벨(관리자 검색어)과 F열 값
    tws = tw["수질"]
    anchor = next(r for r in range(23, tws.max_row + 1)
                  if any("지속삽입가능" in str(tws.cell(r, c).value or "").replace(" ", "") for c in (2, 3)))
    in_tpl = lambda keys: any(str(tws.cell(r, 3).value or "").strip() and any(k.lower() in (str(tws.cell(r, 3).value or "") + str(tws.cell(r, 4).value or "")).replace(" ", "").lower() for k in keys)
                              for r in range(23, tws.max_row + 1))
    missing = [(k, keys) for k, keys in names.items() if wx.get(k) is not None and not in_tpl(keys)]
    labels = {"WaterDepth": "수심"}
    got_rows = []
    r = anchor + 1
    for k, keys in missing:
        while any(tws.cell(r, c).value is not None for c in range(2, 7)):
            r += 1
        got_rows.append((r, ws.cell(r, 3).value, ws.cell(r, 6).value))
        if ws.cell(r, 3).value != labels.get(k, keys[0]) or norm(ws.cell(r, 6).value) != norm(wx.get(k)):
            bad.append(("덧붙임", k, r))
        r += 1
    check(not [b for b in bad if b[0] == "덧붙임"], f"양식에 없는 항목 {[k for k, _ in missing]} → '지속 삽입 가능'({anchor}행) 아래 {[(x[0], x[1], x[2]) for x in got_rows]}")
    appended_rows = {x[0] for x in got_rows}
    tc, oc = cells(tws), cells(ws)
    diff = [k for k in set(tc) | set(oc) if k[1] != 6 and k[0] not in appended_rows and tc.get(k) != oc.get(k)]
    check(not diff, f"라벨(B~D)·예시(E) 보존(단위를 E열에 쓰지 않음) — 다른 칸 {len(diff)}개 {sorted(diff)[:5]}")
    extra = [k for k in oc if k[1] > 6 and not is_formula(oc[k])]
    check(not extra, f"첫 레코드 열(F) 밖 레코드 열에 값 없음 {extra[:5]}")

    # 서식수변 ─────────────────────────────────────────────────────────
    f = file_for(folder, "서식수변")
    print(f"\n=== 서식수변 {os.path.basename(f)}")
    wb, tw = openpyxl.load_workbook(f), openpyxl.load_workbook(os.path.join(tdir, "habitat.xlsx"))
    check(wb.sheetnames == tw.sheetnames, f"시트 = 내장본 {wb.sheetnames} (License 없음)")
    ws = wb["서식 및 수변환경"]
    hx = taxa["HabitatWaterEdge"]

    def side(l, r):
        vals = [v for v in (hx.get(l), hx.get(r)) if v is not None]
        return sum(vals) / len(vals) if vals else None
    items = [hx.get("S1"), side("B2L", "B2R"), hx.get("S3"), hx.get("S4"), side("B5L", "B5R"), side("B6L", "B6R"),
             hx.get("S7"), hx.get("S8"), side("B9L", "B9R"), side("B10L", "B10R")]
    exp = dict(m)
    exp.update({13: hx.get("SurveyUnavailableReason"), 14: hx.get("Note")})
    exp.update({16 + i: v for i, v in enumerate(items)})
    check_column(ws, 6, exp, "F열")
    check(formula_count(wb) == formula_count(tw), f"수식 개수 {formula_count(wb)} = 내장본 {formula_count(tw)}")
    tc, oc = cells(tw["서식 및 수변환경"]), cells(ws)
    diff = [k for k in set(tc) | set(oc) if k[1] < 6 and tc.get(k) != oc.get(k)]
    check(not diff, f"라벨(A~D)·예시(E) 보존 — 다른 칸 {len(diff)}개 {sorted(diff)[:5]}")
    extra = [k for k in oc if k[1] > 6 and not is_formula(oc[k])]
    check(not extra, f"첫 레코드 열(F) 밖 레코드 열에 값 없음 {extra[:5]}")


def main():
    if len(sys.argv) >= 4 and sys.argv[1] == "templates":
        check_templates(sys.argv[2], sys.argv[3])
    elif len(sys.argv) >= 4 and sys.argv[1] == "files":
        tdir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "SDSM_Surveyor_App", "templates", "bulk")
        if "--templates" in sys.argv:
            tdir = sys.argv[sys.argv.index("--templates") + 1]
        check_files(sys.argv[2], sys.argv[3], tdir)
    else:
        print(__doc__)
        sys.exit(2)
    print(f"\n===== {'통과' if not fails else f'실패 {len(fails)}건'} =====")
    sys.exit(0 if not fails else 1)


if __name__ == "__main__":
    main()
